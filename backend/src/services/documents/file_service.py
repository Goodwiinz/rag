"""
File handling and storage service
"""

import hashlib
import logging
import mimetypes
import os
import time
import uuid
import uuid as uuid_module
from pathlib import Path
from typing import Any, BinaryIO, Dict, List, Optional

import aiofiles
import magic
from fastapi import Depends, HTTPException, UploadFile, status
from PIL import Image
from pypdf import PdfReader
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

# Optional pandas import for spreadsheet processing
try:
    import pandas as pd

    PANDAS_AVAILABLE = True
except ImportError:
    pd = None
    PANDAS_AVAILABLE = False

from src.core.config import settings
from src.core.database import get_db
from src.models.document import Document, DocumentType, ProcessingStatus
from src.models.organization import Organization
from src.models.processing import JobPriority, JobStatus, JobType, ProcessingJob
from src.models.user import User, UserRole

logger = logging.getLogger(__name__)


class FileValidationError(Exception):
    """File validation related errors"""

    pass


class FileStorageError(Exception):
    """File storage related errors"""

    pass


class FileService:
    """Service for handling file uploads, validation, and storage"""

    # Map document types to storage bucket names
    TYPE_TO_BUCKET = {
        DocumentType.TEXT: "documents",
        DocumentType.PDF: "documents",
        DocumentType.SPREADSHEET: "documents",
        DocumentType.PRESENTATION: "documents",
        DocumentType.IMAGE: "images",
        DocumentType.AUDIO: "audio",
        DocumentType.VIDEO: "video",
        DocumentType.MULTIMODAL: "documents",
    }

    def __init__(self, db: AsyncSession):
        self.db = db
        self.upload_dir = Path(settings.UPLOAD_DIR)
        self.upload_dir.mkdir(parents=True, exist_ok=True)
        self._storage_backend = settings.STORAGE_BACKEND  # "local", "s3", or "supabase"
        # Legacy compat: SUPABASE_STORAGE_ENABLED=True overrides to "supabase"
        if settings.SUPABASE_STORAGE_ENABLED and self._storage_backend == "local":
            self._storage_backend = "supabase"
        self._storage_helper = None
        self._s3_helper = None

        # Create subdirectories for different file types
        self.create_subdirectories()

    def create_subdirectories(self):
        """Create subdirectories for different file types"""
        subdirs = ["documents", "images", "audio", "video", "temp", "processed"]

        for subdir in subdirs:
            (self.upload_dir / subdir).mkdir(parents=True, exist_ok=True)

    @property
    def storage_helper(self):
        """Lazy-load StorageHelper only when Supabase Storage is enabled."""
        if self._storage_helper is None and self._storage_backend == "supabase":
            from src.core.supabase_client import StorageHelper

            self._storage_helper = StorageHelper()
        return self._storage_helper

    @property
    def s3_helper(self):
        """Lazy-load S3StorageHelper only when S3 backend is active."""
        if self._s3_helper is None and self._storage_backend == "s3":
            from src.core.s3_client import S3StorageHelper

            self._s3_helper = S3StorageHelper()
        return self._s3_helper

    def generate_storage_key(
        self,
        document_type: DocumentType,
        organization_id: str,
        doc_id: str,
        ext: str,
    ) -> tuple[str, str]:
        """Generate a Supabase Storage key for a file.

        Returns (bucket, key) tuple.
        """
        bucket = self.TYPE_TO_BUCKET.get(document_type, "documents")
        key = (
            f"{organization_id}/{doc_id}/{int(time.time())}_{uuid.uuid4().hex[:8]}{ext}"
        )
        return bucket, key

    async def save_file_to_storage(
        self, file: UploadFile, bucket: str, key: str, mime_type: str
    ) -> str:
        """Upload file content to Supabase Storage. Returns the full storage key."""
        content = await file.read()
        file.file.seek(0)
        return self.storage_helper.upload_file(bucket, key, content, mime_type)

    @staticmethod
    def calculate_file_hash_from_bytes(data: bytes) -> str:
        """Calculate SHA-256 hash from raw bytes."""
        return hashlib.sha256(data).hexdigest()

    def get_file_type(self, filename: str, content: bytes = None) -> DocumentType:
        """Determine document type based on file extension and content"""
        # Get MIME type
        mime_type, _ = mimetypes.guess_type(filename)

        # Use python-magic for more accurate detection if content is available
        if content:
            try:
                mime_type = magic.from_buffer(content, mime=True)
            except (OSError, ValueError) as e:
                logger.debug(f"Magic MIME detection failed: {e}")

        # Map MIME types to document types
        if mime_type:
            if mime_type.startswith("text/"):
                return DocumentType.TEXT
            elif mime_type.startswith("image/"):
                return DocumentType.IMAGE
            elif mime_type.startswith("audio/"):
                return DocumentType.AUDIO
            elif mime_type.startswith("video/"):
                return DocumentType.VIDEO
            elif mime_type == "application/pdf":
                return DocumentType.PDF
            elif mime_type in [
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "application/vnd.ms-excel",
            ]:
                return DocumentType.SPREADSHEET
            elif mime_type in [
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                "application/vnd.ms-powerpoint",
            ]:
                return DocumentType.PRESENTATION
            elif mime_type in [
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "application/msword",
            ]:
                return DocumentType.TEXT

        # Fallback to extension-based detection
        ext = Path(filename).suffix.lower()

        if ext in [".txt", ".md", ".rst", ".log"]:
            return DocumentType.TEXT
        elif ext in [".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff", ".webp"]:
            return DocumentType.IMAGE
        elif ext in [".mp3", ".wav", ".ogg", ".flac", ".aac", ".m4a"]:
            return DocumentType.AUDIO
        elif ext in [".mp4", ".avi", ".mkv", ".mov", ".wmv", ".flv", ".webm"]:
            return DocumentType.VIDEO
        elif ext == ".pdf":
            return DocumentType.PDF
        elif ext in [".xlsx", ".xls", ".csv"]:
            return DocumentType.SPREADSHEET
        elif ext in [".pptx", ".ppt"]:
            return DocumentType.PRESENTATION
        elif ext in [".docx", ".doc"]:
            return DocumentType.TEXT

        return DocumentType.MULTIMODAL  # Default for unknown types

    def validate_file(
        self, file: UploadFile, user: User, organization: Organization
    ) -> Dict[str, Any]:
        """Validate uploaded file"""
        # Check file size
        if hasattr(file, "size") and file.size:
            file_size = file.size
        else:
            # Read content to get size if not available
            file.file.seek(0, 2)  # Seek to end
            file_size = file.file.tell()
            file.file.seek(0)  # Reset position

        if file_size > organization.max_file_size_bytes:
            raise FileValidationError(
                f"File size ({file_size} bytes) exceeds maximum allowed size "
                f"({organization.max_file_size_bytes} bytes)"
            )

        # Check storage quota
        if not organization.can_upload_file(file_size):
            raise FileValidationError(
                f"Insufficient storage quota. Available: "
                f"{organization.storage_available_gb:.2f}GB"
            )

        # Check file extension
        allowed_extensions = [
            ".txt",
            ".md",
            ".pdf",
            ".docx",
            ".doc",
            ".xlsx",
            ".xls",
            ".csv",
            ".pptx",
            ".ppt",
            ".jpg",
            ".jpeg",
            ".png",
            ".gif",
            ".bmp",
            ".tiff",
            ".mp3",
            ".wav",
            ".ogg",
            ".flac",
            ".aac",
            ".m4a",
            ".mp4",
            ".avi",
            ".mkv",
            ".mov",
            ".wmv",
            ".flv",
            ".webm",
            ".zip",
            ".rar",
        ]

        file_ext = Path(file.filename).suffix.lower()
        if file_ext not in allowed_extensions:
            raise FileValidationError(f"File extension '{file_ext}' is not allowed")

        # Read content for type detection
        file_content = file.file.read(min(file_size, 8192))  # Read first 8KB
        file.file.seek(0)  # Reset position

        # Determine document type
        document_type = self.get_file_type(file.filename, file_content)

        return {
            "file_size": file_size,
            "document_type": document_type,
            "mime_type": mimetypes.guess_type(file.filename)[0],
        }

    def generate_file_path(
        self, document_type: DocumentType, organization_id: str
    ) -> str:
        """Generate unique file path for uploaded file"""
        # Validate organization_id is a proper UUID to prevent path traversal
        try:
            uuid_module.UUID(organization_id)
        except (ValueError, AttributeError):
            raise ValueError(f"Invalid organization_id format: {organization_id}")

        timestamp = int(time.time())
        unique_id = str(uuid.uuid4())

        type_dir = {
            DocumentType.TEXT: "documents",
            DocumentType.PDF: "documents",
            DocumentType.SPREADSHEET: "documents",
            DocumentType.PRESENTATION: "documents",
            DocumentType.IMAGE: "images",
            DocumentType.AUDIO: "audio",
            DocumentType.VIDEO: "video",
            DocumentType.MULTIMODAL: "documents",
        }

        directory = type_dir.get(document_type, "documents")
        filename = f"{timestamp}_{unique_id}"

        return str(self.upload_dir / directory / f"{organization_id}" / filename)

    async def save_file(self, file: UploadFile, file_path: str) -> str:
        """Save uploaded file to disk"""
        try:
            # Ensure directory exists
            Path(file_path).parent.mkdir(parents=True, exist_ok=True)

            # Save file asynchronously
            async with aiofiles.open(file_path, "wb") as f:
                content = await file.read()
                await f.write(content)

            return file_path

        except Exception as e:
            # Clean up on failure
            if os.path.exists(file_path):
                os.remove(file_path)
            raise FileStorageError(f"Failed to save file: {str(e)}")

    def calculate_file_hash(self, file_path: str) -> str:
        """Calculate SHA-256 hash of file"""
        hash_sha256 = hashlib.sha256()
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(4096), b""):
                hash_sha256.update(chunk)
        return hash_sha256.hexdigest()

    async def upload_file(
        self,
        file: UploadFile,
        title: str,
        user: User,
        organization: Organization,
        tags: List[str] = None,
        is_public: bool = False,
    ) -> Document:
        """Process and store uploaded file"""
        # Track what durably landed so a mid-upload failure can be compensated.
        # The storage object is committed BEFORE the DB rows, and the DB writes
        # span multiple commits — without this, a failure orphans the object,
        # strands a PENDING row, or drifts org storage quota (see the except).
        document = None
        document_committed = False
        quota_committed = False
        processing_job = None
        processing_job_committed = False
        stored_object = None  # (backend, storage_path, file_path) once the object lands
        try:
            # Validate file
            validation_result = self.validate_file(file, user, organization)
            original_ext = Path(file.filename).suffix
            mime_type = validation_result["mime_type"] or "application/octet-stream"

            if self._storage_backend == "s3":
                # --- S3/DO Spaces path ---
                doc_id = str(uuid.uuid4())
                bucket_prefix = self.TYPE_TO_BUCKET.get(
                    validation_result["document_type"], "documents"
                )
                s3_key = (
                    f"{bucket_prefix}/{organization.id}/{doc_id}/"
                    f"{int(time.time())}_{uuid.uuid4().hex[:8]}{original_ext}"
                )

                file_content = await file.read()
                file.file.seek(0)
                file_hash = self.calculate_file_hash_from_bytes(file_content)

                self.s3_helper.upload_file(s3_key, file_content, mime_type)

                document = Document(
                    id=doc_id,
                    title=title,
                    filename=file.filename,
                    file_path=f"s3://{settings.S3_BUCKET_NAME}/{s3_key}",
                    file_size_bytes=validation_result["file_size"],
                    mime_type=mime_type,
                    document_type=validation_result["document_type"],
                    processing_status=ProcessingStatus.PENDING,
                    is_public=is_public,
                    tags=tags or [],
                    organization_id=organization.id,
                    uploaded_by_user_id=user.id,
                    storage_path=s3_key,
                    storage_backend="s3",
                )

            elif self._storage_backend == "supabase":
                # --- Supabase Storage path ---
                doc_id = str(uuid.uuid4())
                bucket, key = self.generate_storage_key(
                    validation_result["document_type"],
                    str(organization.id),
                    doc_id,
                    original_ext,
                )

                file_content = await file.read()
                file.file.seek(0)
                file_hash = self.calculate_file_hash_from_bytes(file_content)

                storage_key = self.storage_helper.upload_file(
                    bucket, key, file_content, mime_type
                )

                document = Document(
                    id=doc_id,
                    title=title,
                    filename=file.filename,
                    file_path=f"supabase://{storage_key}",
                    file_size_bytes=validation_result["file_size"],
                    mime_type=mime_type,
                    document_type=validation_result["document_type"],
                    processing_status=ProcessingStatus.PENDING,
                    is_public=is_public,
                    tags=tags or [],
                    organization_id=organization.id,
                    uploaded_by_user_id=user.id,
                    storage_path=storage_key,
                    storage_backend="supabase",
                )

            else:
                # --- Local filesystem path (default) ---
                file_path = self.generate_file_path(
                    validation_result["document_type"], str(organization.id)
                )
                file_path_with_ext = f"{file_path}{original_ext}"
                saved_path = await self.save_file(file, file_path_with_ext)
                file_hash = self.calculate_file_hash(saved_path)

                document = Document(
                    title=title,
                    filename=file.filename,
                    file_path=saved_path,
                    file_size_bytes=validation_result["file_size"],
                    mime_type=mime_type,
                    document_type=validation_result["document_type"],
                    processing_status=ProcessingStatus.PENDING,
                    is_public=is_public,
                    tags=tags or [],
                    organization_id=organization.id,
                    uploaded_by_user_id=user.id,
                    storage_backend="local",
                )

            # Capture the storage identity as plain values now, before any commit
            # or rollback can expire the ORM instance. The rollback path must be
            # able to delete the object without reading the (possibly expired)
            # instance — a sync attribute read on an expired instance raises
            # MissingGreenlet inside an async session.
            stored_object = (
                document.storage_backend,
                document.storage_path,
                document.file_path,
            )

            # Add file hash as metadata
            document.add_metadata("file_hash", file_hash)
            document.add_metadata("original_filename", file.filename)

            self.db.add(document)
            await self.db.commit()
            await self.db.refresh(document)
            document_committed = True

            # Atomically update organization storage usage
            await self.db.execute(
                Organization.storage_usage_update(
                    organization.id, validation_result["file_size"]
                )
            )
            await self.db.commit()
            quota_committed = True

            # Create processing job for document ingestion
            processing_job = ProcessingJob(
                job_type=JobType.DOCUMENT_INGESTION,
                status=JobStatus.PENDING,
                priority=JobPriority.NORMAL,
                document_id=document.id,
                organization_id=organization.id,
                created_by_user_id=user.id,
                parameters={
                    "document_id": str(document.id),
                    "file_path": document.effective_file_path,
                    "document_type": document.document_type.value,
                    "mime_type": document.mime_type,
                    "storage_backend": document.storage_backend,
                },
                config={"max_retries": 3, "timeout_seconds": 300},
                total_steps=5,
                queue_name="document_processing",
            )

            self.db.add(processing_job)
            await self.db.commit()
            await self.db.refresh(processing_job)
            processing_job_committed = True

            # Queue the job for processing
            from src.tasks.processing_tasks import process_document_ingestion

            process_document_ingestion.delay(str(processing_job.id))

            return document

        except Exception as e:
            await self.db.rollback()
            # Compensate whatever durably landed before the failure. First reverse
            # the DB (soft-delete the row + revert quota + drop the stray job in one
            # commit); only if that succeeds do we delete the storage object. If the
            # DB reversal itself fails (e.g. the connection is still bad), the row
            # stays live, so we keep the object as a sweepable orphan rather than
            # orphan a live row from its backing file. `reversal_ok` starts True when
            # nothing was committed (no live row to protect).
            reversal_ok = not document_committed
            if document_committed and document is not None:
                try:
                    # A committed ProcessingJob only exists when the enqueue
                    # (.delay) failed — drop it so it can't run against the
                    # soft-deleted document.
                    if processing_job_committed and processing_job is not None:
                        await self.db.delete(processing_job)
                    document.soft_delete()
                    if quota_committed:
                        await self.db.execute(
                            Organization.storage_usage_update(
                                organization.id, -validation_result["file_size"]
                            )
                        )
                    await self.db.commit()
                    reversal_ok = True
                except Exception:
                    await self.db.rollback()
                    logger.warning(
                        "upload rollback: failed to reverse committed row/quota for "
                        "document %s; leaving storage object as a sweepable orphan",
                        getattr(document, "id", None),
                        exc_info=True,
                    )
            # Delete by captured primitives (never the possibly-expired instance).
            if reversal_ok and stored_object is not None:
                backend, storage_path, obj_file_path = stored_object
                try:
                    self._delete_stored_object(backend, storage_path, obj_file_path)
                except Exception:
                    logger.warning(
                        "upload rollback: orphaned storage object %s (recoverable "
                        "by a later sweep)",
                        storage_path or obj_file_path,
                        exc_info=True,
                    )
            raise FileStorageError(f"Failed to upload file: {str(e)}")

    def extract_text_content(self, document: Document) -> str:
        """Extract text content from document.

        Uses local_file_for_document to transparently handle Supabase-backed files.
        """
        from src.services.documents.storage_utils import local_file_for_document

        try:
            with local_file_for_document(document) as file_path:
                return self._extract_text_from_path(file_path, document)
        except Exception:
            # Re-raise instead of returning an "Error extracting text: ..."
            # string — callers stored that string as document.content_text and
            # marked the document COMPLETED/indexed, indexing an error message
            # as the document's content. The processing pipeline (and agent
            # tools) catch this and fail the document / return a clean error.
            logger.warning(
                "Text extraction failed for document %s",
                getattr(document, "id", "?"),
                exc_info=True,
            )
            raise

    def _extract_text_from_path(self, file_path: str, document: Document) -> str:
        """Extract text from a local file path."""
        try:
            if document.document_type == DocumentType.TEXT:
                # Simple text file
                with open(file_path, "r", encoding="utf-8") as f:
                    return f.read()

            elif document.document_type == DocumentType.PDF:
                # PDF file
                text = []
                with open(file_path, "rb") as file:
                    pdf_reader = PdfReader(file)
                    for page in pdf_reader.pages:
                        # `or ""` so a single page that yields None doesn't
                        # TypeError the whole join (losing every other page).
                        text.append(page.extract_text() or "")
                return "\n".join(text)

            elif document.document_type in [DocumentType.SPREADSHEET]:
                # Excel file - process if pandas is available
                if not PANDAS_AVAILABLE:
                    logger.warning(
                        "Pandas not available, skipping spreadsheet processing"
                    )
                    # Empty (no text), not a capability-gap string — the latter
                    # would be indexed as the document's content.
                    return ""
                try:
                    if file_path.endswith(".csv"):
                        df = pd.read_csv(file_path)
                    else:
                        df = pd.read_excel(file_path)
                    return df.to_string()
                except Exception as e:
                    # pandas IS available here (the missing-pandas capability
                    # gap returned "" above), so this is a real parse failure of
                    # an existing file. Raise — returning "" would index the
                    # corrupt file as empty and mark it COMPLETED/indexed, i.e.
                    # silently unqueryable. The pipeline catches this → FAILED.
                    logger.error(f"Spreadsheet extraction failed for {file_path}: {e}")
                    raise

            elif document.document_type == DocumentType.PRESENTATION:
                # PowerPoint file (basic extraction)
                try:
                    from pptx import Presentation

                    prs = Presentation(file_path)
                    text = []
                    for slide in prs.slides:
                        for shape in slide.shapes:
                            if hasattr(shape, "text"):
                                text.append(shape.text)
                    return "\n".join(text)
                except ImportError:
                    # python-pptx not installed: a capability gap, not a bad
                    # file. Empty (like the missing-pandas case) — don't fail the
                    # document over a missing optional dependency.
                    logger.warning(
                        "python-pptx not available, skipping presentation processing"
                    )
                    return ""
                except (ValueError, IOError) as e:
                    # Corrupt/unreadable presentation — raise so the document is
                    # marked FAILED rather than indexed empty + COMPLETED.
                    logger.error(f"Presentation extraction failed for {file_path}: {e}")
                    raise

            elif document.document_type == DocumentType.IMAGE:
                # For images, we could use OCR here (integrate with Tesseract)
                return ""

            else:
                return ""

        except Exception:
            # Genuine extraction failure must fail the document, not be indexed
            # as an "Error extracting text: ..." content string. See
            # extract_text_content above.
            raise

    def extract_metadata(self, document: Document) -> Dict[str, Any]:
        """Extract metadata from document.

        Uses local_file_for_document to transparently handle Supabase-backed files.
        """
        from src.services.documents.storage_utils import local_file_for_document

        with local_file_for_document(document) as file_path:
            return self._extract_metadata_from_path(file_path, document)

    def _extract_metadata_from_path(
        self, file_path: str, document: Document
    ) -> Dict[str, Any]:
        """Extract metadata from a local file path."""
        metadata = {}

        try:
            if document.document_type == DocumentType.IMAGE:
                # Image metadata
                with Image.open(file_path) as img:
                    metadata.update(
                        {
                            "width": img.width,
                            "height": img.height,
                            "format": img.format,
                            "mode": img.mode,
                        }
                    )

            elif document.document_type == DocumentType.PDF:
                # PDF metadata
                with open(file_path, "rb") as file:
                    pdf_reader = PdfReader(file)
                    if pdf_reader.metadata:
                        metadata.update(
                            {
                                "title": pdf_reader.metadata.get("/Title", ""),
                                "author": pdf_reader.metadata.get("/Author", ""),
                                "subject": pdf_reader.metadata.get("/Subject", ""),
                                "creator": pdf_reader.metadata.get("/Creator", ""),
                                "producer": pdf_reader.metadata.get("/Producer", ""),
                                "creation_date": str(
                                    pdf_reader.metadata.get("/CreationDate", "")
                                ),
                                "modification_date": str(
                                    pdf_reader.metadata.get("/ModDate", "")
                                ),
                            }
                        )
                    metadata["page_count"] = len(pdf_reader.pages)

            elif document.document_type in [
                DocumentType.SPREADSHEET,
                DocumentType.PRESENTATION,
            ]:
                # Excel/PowerPoint metadata (basic)
                file_stat = os.stat(file_path)
                metadata.update(
                    {
                        "size": file_stat.st_size,
                        "created": file_stat.st_ctime,
                        "modified": file_stat.st_mtime,
                    }
                )

        except Exception as e:
            metadata["extraction_error"] = str(e)

        return metadata

    def _delete_stored_object(
        self,
        storage_backend: Optional[str],
        storage_path: Optional[str],
        file_path: Optional[str],
    ) -> None:
        """Delete a stored object by its raw identity (backend / key / path).

        Takes plain values rather than a Document so it can run in the upload
        rollback path, where the ORM instance may be expired — a sync attribute
        read on an expired instance raises MissingGreenlet in an async session.
        """
        if storage_backend == "s3" and storage_path:
            from src.core.s3_client import S3StorageHelper

            S3StorageHelper().delete_file(storage_path)
        elif storage_backend == "supabase" and storage_path:
            from src.core.supabase_client import parse_storage_key

            bucket, key = parse_storage_key(storage_path)
            self.storage_helper.delete_file(bucket, key)
        elif file_path and os.path.exists(file_path):
            os.remove(file_path)

    def delete_physical_file(self, document: Document) -> None:
        """Delete every stored object a document owns.

        Three object classes accumulate for one document:

        1. the original upload (``storage_path``), honoring ``storage_backend``
           (s3 / supabase / local disk);
        2. the canonical DO-KB text mirror ``documents/{org}/{doc}.txt``; and
        3. figure PNG crops under ``figures/{org}/{doc}/``.

        Classes 2 and 3 are always written to S3/Spaces by the KB-ingest and
        figure-extraction services (they upload via ``S3StorageHelper`` directly,
        independent of the document's ``storage_backend``), so they are removed
        from S3 whenever a client is available. Deleting only ``storage_path``
        left them behind as retained user content after delete (audit finding
        D6). Auxiliary cleanup is best-effort and never raises; the original
        delete keeps its raise-on-failure contract so callers still log it as a
        recoverable orphan.
        """
        original_error: Optional[Exception] = None
        try:
            self._delete_stored_object(
                document.storage_backend, document.storage_path, document.file_path
            )
        except Exception as exc:  # re-raised below, AFTER aux cleanup runs
            original_error = exc

        # Run the derived-object cleanup regardless of the original delete's
        # outcome, and never let it mask or replace the original error.
        self._delete_auxiliary_objects(document)

        if original_error is not None:
            raise original_error

    @staticmethod
    def _s3_helper_or_none():
        """An ``S3StorageHelper`` when S3/Spaces is configured, else ``None``.

        The canonical text mirror and figure crops live on S3 regardless of the
        document's ``storage_backend``, so their cleanup targets S3 whenever it
        is available — NOT ``self.s3_helper``, which is ``None`` unless the
        service's own backend is s3. A missing S3 config is not an error (those
        object classes were never created without S3), so this returns ``None``
        instead of raising (``S3StorageHelper.__init__`` raises ``RuntimeError``
        when unconfigured).
        """
        try:
            from src.core.s3_client import S3StorageHelper

            return S3StorageHelper()
        except Exception:
            return None

    def _delete_auxiliary_objects(self, document: Document) -> None:
        """Best-effort delete of a document's canonical KB text mirror + figure
        PNG crops from S3/Spaces.

        Every key is derived strictly from the document row (its org id + id),
        so this can only ever touch objects that belong to THIS document — never
        an object a live document references. Never raises: each storage error
        is logged and swallowed (an undeleted auxiliary object is a sweepable
        orphan, surfaced by ``storage_reconcile``).
        """
        helper = self._s3_helper_or_none()
        if helper is None:
            return  # No S3 → these object classes never existed.

        # Key derivations live in the dependency-light ``object_keys`` module, so
        # this hot path derives them WITHOUT importing the heavy KB-ingest /
        # figure-extraction (PDF/ML) packages, and can't drift from the writers.
        from src.services.documents.object_keys import (
            canonical_text_key,
            figure_object_prefix,
        )

        doc_id = getattr(document, "id", "?")

        # (2) Canonical DO-KB text mirror: documents/{org}/{doc}.txt.
        try:
            helper.delete_file(canonical_text_key(document))
        except Exception:
            logger.warning(
                "Failed to delete canonical KB text object for document %s "
                "(orphan, recoverable by storage_reconcile)",
                doc_id,
                exc_info=True,
            )

        # (3) Figure PNG crops: figures/{org}/{doc}/... — list the doc-scoped
        #     prefix (it embeds the doc UUID, so it matches ONLY this document's
        #     figures) and delete each object.
        try:
            for key in helper.list_objects(figure_object_prefix(document)):
                helper.delete_file(key)
        except Exception:
            logger.warning(
                "Failed to enumerate/delete figure objects for document %s "
                "(orphan, recoverable by storage_reconcile)",
                doc_id,
                exc_info=True,
            )

    async def delete_file(self, document: Document, user: User) -> bool:
        """Delete file and update storage.

        Mirrors the cascade of ``documents.delete_document``: this is the second
        live delete surface for the same rows, and previously it soft-deleted
        only the Document + quota — leaving the document's Entity rows live in
        Postgres and its subgraph / DO KB data source orphaned forever (a
        deleted doc that still surfaces in retrieval and graph results). Reap
        all satellites the same way.
        """
        # Check permissions first — outside the try so a 403 propagates as-is
        # instead of being wrapped into a FileStorageError.
        if document.uploaded_by_user_id != user.id and not user.has_permission(
            UserRole.ADMIN
        ):
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Can only delete your own files or require admin role",
            )

        # Capture ids before soft_delete / commit for the post-commit satellite
        # cleanup (the ORM object's attributes stay readable, but be explicit).
        document_id = str(document.id)
        organization_id = str(document.organization_id)

        # Make the DB the source of truth FIRST: soft-delete the record + its
        # Entity rows + processing jobs and revert quota, then commit. Only after
        # that succeeds do we remove the physical object. Deleting the object
        # before the commit meant a commit failure rolled back the row while the
        # storage object was already irreversibly gone — a live row pointing at a
        # missing file (every later download / content / reprocess 403/404, quota
        # still counted). With this order a failure leaves at worst a sweepable
        # orphan object, never a live row whose backing file is gone.
        try:
            from datetime import datetime

            from src.models.entity import Entity

            await self.db.execute(
                update(Entity)
                .where(Entity.document_id == document.id, Entity.is_deleted == False)
                .values(is_deleted=True, deleted_at=datetime.utcnow())
            )
            await self.db.execute(
                update(ProcessingJob)
                .where(
                    ProcessingJob.document_id == document.id,
                    ProcessingJob.is_deleted == False,
                )
                .values(is_deleted=True, deleted_at=datetime.utcnow())
            )

            document.soft_delete()
            await self.db.execute(
                Organization.storage_usage_update(
                    document.organization_id, -document.file_size_bytes
                )
            )
            await self.db.commit()
        except Exception as e:
            await self.db.rollback()
            raise FileStorageError(f"Failed to delete file: {str(e)}")

        # Best-effort physical delete AFTER the commit. A failure here must NOT
        # roll back the committed soft-delete — log it as a recoverable orphan.
        try:
            self.delete_physical_file(document)
        except Exception as delete_error:
            logger.warning(
                "Soft-deleted document %s but failed to remove its storage "
                "object (orphan, recoverable by a later sweep): %s",
                document.id,
                delete_error,
                exc_info=True,
            )

        # Reap the DO KB data source + Neo4j subgraph (best-effort, post-commit,
        # never blocks the delete) — same as documents.delete_document.
        await self._cleanup_satellites_on_delete(document, document_id, organization_id)

        return True

    async def _cleanup_satellites_on_delete(
        self, document: Document, document_id: str, organization_id: str
    ) -> None:
        """Best-effort removal of a deleted document's DO KB data source and
        Neo4j subgraph. Each step is failure-isolated: a KB/Neo4j outage during
        delete must never fail or block the user's delete (the Postgres rows are
        already gone). Failures are logged as recoverable drift."""
        # DO KB: unsync so the deleted doc stops surfacing in retrieval and stops
        # leaking storage. unsync_document_from_kb no-ops when DO_KB is off or the
        # doc has no data source, and never raises.
        try:
            if document.do_kb_data_source_uuid:
                from src.services.do_kb import unsync_document_from_kb

                await unsync_document_from_kb(self.db, document)
        except Exception:  # noqa: BLE001
            logger.warning(
                "do_kb cleanup on file delete failed",
                extra={"document_id": document_id},
                exc_info=True,
            )

        # Neo4j: reap this document's relationships then its now-orphaned entity
        # nodes (shared across docs — no blind DETACH DELETE), org-scoped. The KG
        # service is synchronous, so offload to a worker thread.
        try:
            import asyncio

            from src.services.knowledge_graph.knowledge_graph_service import (
                KnowledgeGraphService,
            )

            await asyncio.to_thread(
                lambda: KnowledgeGraphService().delete_document_graph(
                    document_id, organization_id
                )
            )
        except Exception:  # noqa: BLE001
            logger.warning(
                "knowledge-graph cleanup on file delete failed",
                extra={"document_id": document_id},
                exc_info=True,
            )

    async def get_file_stats(self, organization_id: str) -> Dict[str, Any]:
        """Get file statistics for organization"""
        from sqlalchemy import func

        # Total files by type
        files_by_type_stmt = (
            select(
                Document.document_type,
                func.count(Document.id).label("count"),
                func.sum(Document.file_size_bytes).label("total_size"),
            )
            .where(
                Document.organization_id == organization_id,
                Document.is_deleted == False,
            )
            .group_by(Document.document_type)
        )
        files_by_type_result = await self.db.execute(files_by_type_stmt)
        files_by_type = files_by_type_result.all()

        # Processing status distribution
        processing_stats_stmt = (
            select(Document.processing_status, func.count(Document.id).label("count"))
            .where(
                Document.organization_id == organization_id,
                Document.is_deleted == False,
            )
            .group_by(Document.processing_status)
        )
        processing_stats_result = await self.db.execute(processing_stats_stmt)
        processing_stats = processing_stats_result.all()

        return {
            "files_by_type": [
                {
                    "type": stat.document_type.value,
                    "count": stat.count,
                    "total_size_mb": (
                        stat.total_size / (1024 * 1024) if stat.total_size else 0
                    ),
                }
                for stat in files_by_type
            ],
            "processing_stats": [
                {"status": stat.processing_status.value, "count": stat.count}
                for stat in processing_stats
            ],
        }


def get_file_service(db: AsyncSession = Depends(get_db)) -> FileService:
    """Get file service instance"""
    return FileService(db)
